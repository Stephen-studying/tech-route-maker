"""Render editable PowerPoint shapes from a technical-route JSON model."""

import os
import sys
import zipfile
from datetime import datetime, timezone
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Pt

from route_common import (
    build_layout,
    edge_is_feedback,
    edge_points,
    get_style,
    load_route,
    show_node_edges,
    stage_flow_segments,
    wrap_text,
)


LANDSCAPE_W = 12192000
LANDSCAPE_H = 6858000
PORTRAIT_W = 6858000
PORTRAIT_H = 12192000
FONT_NAME = "Microsoft YaHei"
FIXED_TIME = datetime(2026, 1, 1, tzinfo=timezone.utc)


def rgb(value):
    return RGBColor.from_string(str(value or "#000000").replace("#", "").upper())


def palette_item(style, index):
    palette = style.get("palette") or [style["stage_fill"]]
    return palette[index % len(palette)]


def slide_size_for(layout):
    if layout["orientation"] == "research-framework-template":
        width = 9144000
        return width, int(width * layout["height"] / layout["width"])
    if layout["orientation"] == "vertical" and str(layout.get("layout_name", "")).startswith("cn-"):
        return PORTRAIT_W, PORTRAIT_H
    if layout["height"] > layout["width"] * 1.18:
        return PORTRAIT_W, PORTRAIT_H
    return LANDSCAPE_W, LANDSCAPE_H


def transform_box(box, scale, ox, oy):
    return {
        "x": int(ox + box["x"] * scale),
        "y": int(oy + box["y"] * scale),
        "w": int(box["w"] * scale),
        "h": int(box["h"] * scale),
    }


def set_east_asian_font(run):
    rpr = run._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea"):
        existing = rpr.find(tag, rpr.nsmap)
        if existing is None:
            existing = OxmlElement(tag)
            rpr.append(existing)
        existing.set("typeface", FONT_NAME)


def wrapped_lines(text, max_chars, max_lines):
    lines = []
    for segment in str(text or "").splitlines() or [""]:
        lines.extend(wrap_text(segment, max_chars, max_lines))
    return lines[:max_lines] or [""]


def remove_theme_style(shape):
    style = shape._element.find(qn("p:style"))
    if style is not None:
        shape._element.remove(style)


def set_shape_text(
    shape,
    text,
    font_size,
    font_color,
    bold=False,
    max_chars=18,
    max_lines=4,
    word_wrap=True,
):
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = word_wrap
    frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    frame.margin_left = Pt(3)
    frame.margin_right = Pt(3)
    frame.margin_top = Pt(2)
    frame.margin_bottom = Pt(2)
    lines = wrapped_lines(text, max_chars, max_lines)
    for index, line in enumerate(lines):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.alignment = PP_ALIGN.CENTER
        paragraph.space_before = Pt(0)
        paragraph.space_after = Pt(0)
        run = paragraph.add_run()
        run.text = line
        run.font.name = FONT_NAME
        run.font.size = Pt(font_size)
        run.font.bold = bold
        run.font.color.rgb = rgb(font_color)
        set_east_asian_font(run)


def add_shape(
    slide,
    shape_type,
    box,
    text,
    fill,
    stroke,
    font_color,
    font_size=12,
    bold=False,
    dash=False,
    max_chars=18,
    max_lines=4,
    line_width=1.0,
    name=None,
    word_wrap=True,
):
    shape = slide.shapes.add_shape(shape_type, box["x"], box["y"], box["w"], box["h"])
    if name:
        shape.name = name
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    shape.line.color.rgb = rgb(stroke)
    shape.line.width = Pt(line_width)
    if dash:
        shape.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    remove_theme_style(shape)
    set_shape_text(
        shape,
        text,
        font_size,
        font_color,
        bold,
        max_chars,
        max_lines,
        word_wrap,
    )
    return shape


def add_line(slide, x1, y1, x2, y2, stroke, arrow=True, dash=False, width=1.5, name=None):
    connector = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT,
        int(x1),
        int(y1),
        int(x2),
        int(y2),
    )
    if name:
        connector.name = name
    connector.line.color.rgb = rgb(stroke)
    connector.line.width = Pt(width)
    if dash:
        connector.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    remove_theme_style(connector)
    if arrow:
        line_xml = connector.line._get_or_add_ln()
        for existing in list(line_xml):
            if existing.tag.endswith("}tailEnd"):
                line_xml.remove(existing)
        tail = OxmlElement("a:tailEnd")
        tail.set("type", "triangle")
        line_xml.append(tail)
    return connector


def add_template_shapes(slide, layout, route, style, scale, ox, oy):
    for header in layout.get("headers") or []:
        box = transform_box(header, scale, ox, oy)
        if header["role"] == "content":
            fill = style["content_header_fill"]
            stroke = style["content_header_stroke"]
        else:
            fill = style["side_header_fill"]
            stroke = style["side_header_stroke"]
        add_shape(
            slide,
            MSO_SHAPE.RECTANGLE,
            box,
            header["text"],
            fill,
            stroke,
            style["text"],
            17.5,
            True,
            True,
            12,
            2,
            1.2,
            f"Template Header {header['id']}",
        )

    group_strokes = style.get("group_strokes") or [style["stage_stroke"]]
    for stage in layout["stages"]:
        group_stroke = group_strokes[stage["group_stroke_index"] % len(group_strokes)]
        add_shape(
            slide,
            MSO_SHAPE.RECTANGLE,
            transform_box(stage, scale, ox, oy),
            "",
            palette_item(style, stage["index"]),
            group_stroke,
            style["text"],
            9,
            False,
            True,
            18,
            1,
            1.2,
            f"Content Group {stage['content_label']}",
        )
        label = transform_box(stage["content_label_box"], scale, ox, oy)
        stacked = "\n".join(str(stage["content_label"]).replace("\n", "").replace(" ", ""))
        add_shape(
            slide,
            MSO_SHAPE.RECTANGLE,
            label,
            stacked,
            "#FFFFFF",
            style["template_border"],
            style["text"],
            12,
            True,
            True,
            1,
            12,
            1.0,
            f"Content Label {stage['content_label']}",
        )
        add_shape(
            slide,
            MSO_SHAPE.OVAL,
            transform_box(stage["logic_box"], scale, ox, oy),
            stage["logic_label"],
            style["logic_fill"],
            style["logic_stroke"],
            style["text"],
            12.5,
            True,
            False,
            9,
            3,
            1.2,
            f"Logic {stage['logic_label']}",
            False,
        )
        add_shape(
            slide,
            MSO_SHAPE.OVAL,
            transform_box(stage["method_box"], scale, ox, oy),
            stage["method_label"],
            style["method_fill"],
            style["method_stroke"],
            style["text"],
            11.0,
            False,
            False,
            9,
            3,
            1.2,
            f"Method {stage['method_label']}",
            False,
        )

    for index, arrow in enumerate(layout.get("side_arrows") or [], start=1):
        shape_type = MSO_SHAPE.LEFT_ARROW if arrow["direction"] == "left" else MSO_SHAPE.RIGHT_ARROW
        add_shape(
            slide,
            shape_type,
            transform_box(arrow, scale, ox, oy),
            "",
            "#FFFFFF",
            style["template_arrow"],
            style["text"],
            8,
            False,
            False,
            1,
            1,
            1.1,
            f"Template Side Arrow {index}",
        )
    for index, segment in enumerate(layout.get("spine_segments") or [], start=1):
        add_line(
            slide,
            ox + segment["x1"] * scale,
            oy + segment["y1"] * scale,
            ox + segment["x2"] * scale,
            oy + segment["y2"] * scale,
            style["template_spine"],
            True,
            False,
            4.0,
            f"Logic Spine {index}",
        )

    nodes = {node_id: transform_box(box, scale, ox, oy) for node_id, box in layout["nodes"].items()}
    for stage in route["stages"]:
        for node in stage["nodes"]:
            if node["id"] not in nodes:
                continue
            add_shape(
                slide,
                MSO_SHAPE.RECTANGLE,
                nodes[node["id"]],
                node["label"],
                style["node_fill"],
                style["template_border"],
                style["text"],
                14.5,
                True,
                True,
                30,
                2,
                1.0,
                f"Content Row {node['id']}",
            )


def add_generic_shapes(slide, layout, route, style, scale, ox, oy):
    slide_w, _slide_h = slide.part.package.presentation_part.presentation.slide_width, None
    title_w = int(min(layout["width"] - 120, max(520, layout["width"] * 0.78)) * scale)
    title_x = int((slide_w - title_w) / 2)
    title_y = int(oy + 18 * scale)
    title_h = int(58 * scale)
    title_fill = style["background"] if layout["orientation"] in {"mainline", "a4stage"} else palette_item(style, 0)
    title_stroke = title_fill if layout["orientation"] in {"mainline", "a4stage"} else style["stage_stroke"]
    title_size = 29 if layout["orientation"] == "mainline" else 25 if layout["orientation"] == "a4stage" else 21
    add_shape(
        slide,
        MSO_SHAPE.RECTANGLE if layout["orientation"] in {"mainline", "a4stage"} else MSO_SHAPE.ROUNDED_RECTANGLE,
        {"x": title_x, "y": title_y, "w": title_w, "h": title_h},
        route["title"],
        title_fill,
        title_stroke,
        style["text"],
        title_size,
        True,
        False,
        34,
        2,
        0.8,
        "Title",
    )
    if route.get("subtitle"):
        add_shape(
            slide,
            MSO_SHAPE.RECTANGLE,
            {
                "x": title_x,
                "y": int(title_y + title_h + 5 * scale),
                "w": title_w,
                "h": int(28 * scale),
            },
            route["subtitle"],
            style["background"],
            style["background"],
            style["muted"],
            10.5,
            False,
            False,
            48,
            1,
            0.1,
            "Subtitle",
        )

    if layout["orientation"] == "vertical" and layout["stages"]:
        axis_x = ox + layout["stages"][0].get("axis_x", 100) * scale
        first = layout["stages"][0]
        last = layout["stages"][-1]
        add_line(
            slide,
            axis_x,
            oy + (first["y"] + 10) * scale,
            axis_x,
            oy + (last["y"] + last["h"] - 10) * scale,
            style["accent"],
            False,
            False,
            2.0,
            "Vertical Phase Axis",
        )

    for stage in layout["stages"]:
        box = transform_box(stage, scale, ox, oy)
        stage_shape = MSO_SHAPE.ROUNDED_RECTANGLE
        add_shape(
            slide,
            stage_shape,
            box,
            "",
            palette_item(style, stage["index"]),
            style["stage_stroke"],
            style["text"],
            9,
            False,
            bool(style.get("stage_dash")),
            18,
            1,
            0.9,
            f"Stage {stage['title']}",
        )
        if stage.get("layout") == "mainline":
            main = transform_box(stage["main_box"], scale, ox, oy)
            add_shape(
                slide,
                MSO_SHAPE.ROUNDED_RECTANGLE,
                main,
                stage["title"],
                style.get("main_fill", "#DBEAFE"),
                style.get("main_stroke", style["header_fill"]),
                style.get("main_text", style["text"]),
                15.5,
                True,
                False,
                14,
                2,
                1.2,
                f"Main {stage['title']}",
            )
        elif stage.get("layout") == "a4stage":
            label = f"{stage['index'] + 1:02d}  {stage['title']}"
            label_w = min(box["w"] - int(72 * scale), max(int(340 * scale), len(label) * int(20 * scale)))
            add_shape(
                slide,
                MSO_SHAPE.ROUNDED_RECTANGLE,
                {
                    "x": int(box["x"] + (box["w"] - label_w) / 2),
                    "y": int(box["y"] + 20 * scale),
                    "w": int(label_w),
                    "h": int(48 * scale),
                },
                label,
                style["header_fill"],
                style["header_fill"],
                style["header_text"],
                14.5,
                True,
                False,
                16,
                1,
                0.8,
                f"Header {stage['title']}",
            )
        elif stage.get("label_box"):
            label = transform_box(stage["label_box"], scale, ox, oy)
            add_shape(
                slide,
                MSO_SHAPE.ROUNDED_RECTANGLE,
                label,
                stage["title"],
                style["header_fill"],
                style["header_fill"],
                style["header_text"],
                10.5,
                True,
                False,
                12,
                3,
                0.8,
                f"Label {stage['title']}",
            )
        elif layout["orientation"] == "vertical":
            label_w = min(box["w"] * 0.34, int(230 * scale))
            add_shape(
                slide,
                MSO_SHAPE.ROUNDED_RECTANGLE,
                {
                    "x": int(box["x"] + 28 * scale),
                    "y": int(box["y"] - 17 * scale),
                    "w": int(label_w),
                    "h": int(34 * scale),
                },
                stage["title"],
                style["header_fill"],
                style["header_fill"],
                style["header_text"],
                11,
                True,
                False,
                14,
                2,
                0.8,
                f"Header {stage['title']}",
            )

    scaled_nodes = {
        node_id: transform_box(box, scale, ox, oy)
        for node_id, box in layout["nodes"].items()
    }
    if show_node_edges(route):
        for edge in route["edges"]:
            if not edge.get("valid", True) or edge["from"] not in scaled_nodes or edge["to"] not in scaled_nodes:
                continue
            source = scaled_nodes[edge["from"]]
            target = scaled_nodes[edge["to"]]
            points = edge_points(source, target, layout["orientation"])
            dashed = edge_is_feedback(edge, source, target)
            for index in range(len(points) - 1):
                x1, y1 = points[index]
                x2, y2 = points[index + 1]
                add_line(
                    slide,
                    x1,
                    y1,
                    x2,
                    y2,
                    style["line"],
                    index == len(points) - 2,
                    dashed,
                    1.4,
                )
    else:
        for index, (x1, y1, x2, y2) in enumerate(stage_flow_segments(layout), start=1):
            add_line(
                slide,
                ox + x1 * scale,
                oy + y1 * scale,
                ox + x2 * scale,
                oy + y2 * scale,
                style["line"],
                True,
                False,
                1.5,
                f"Stage Flow {index}",
            )

    for stage in route["stages"]:
        for node in stage["nodes"]:
            if node["id"] not in scaled_nodes:
                continue
            node_size = 13.2 if layout["orientation"] == "mainline" else 12.5 if layout["orientation"] == "a4stage" else 9.5
            add_shape(
                slide,
                MSO_SHAPE.ROUNDED_RECTANGLE,
                scaled_nodes[node["id"]],
                node["label"],
                style["node_fill"],
                style["node_stroke"],
                style["text"],
                node_size,
                False,
                False,
                18,
                3,
                1.0,
                f"Node {node['id']}",
            )

    bar = layout.get("output_bar")
    if bar and bar.get("text"):
        add_shape(
            slide,
            MSO_SHAPE.ROUNDED_RECTANGLE,
            transform_box(bar, scale, ox, oy),
            "成果输出：" + str(bar["text"]),
            style.get("output_fill", "#E0F2FE"),
            style.get("output_stroke", style["header_fill"]),
            style["text"],
            14,
            True,
            False,
            48,
            2,
            1.2,
            "Output",
        )


def build_presentation(route):
    layout = build_layout(route)
    route = layout["route"]
    style = get_style(route)
    slide_w, slide_h = slide_size_for(layout)
    presentation = Presentation()
    presentation.slide_width = slide_w
    presentation.slide_height = slide_h
    presentation.core_properties.author = "tech-route-maker"
    presentation.core_properties.last_modified_by = "tech-route-maker"
    presentation.core_properties.created = FIXED_TIME
    presentation.core_properties.modified = FIXED_TIME
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    scale = min((slide_w - 260000) / layout["width"], (slide_h - 260000) / layout["height"])
    ox = (slide_w - layout["width"] * scale) / 2
    oy = (slide_h - layout["height"] * scale) / 2
    if layout["orientation"] == "research-framework-template":
        add_template_shapes(slide, layout, route, style, scale, ox, oy)
    else:
        add_generic_shapes(slide, layout, route, style, scale, ox, oy)
    return presentation


def normalize_zip(path):
    path = Path(path)
    temp = path.with_suffix(path.suffix + ".tmp")
    with zipfile.ZipFile(path, "r") as source:
        members = [(item, source.read(item.filename)) for item in source.infolist()]
    with zipfile.ZipFile(temp, "w", zipfile.ZIP_DEFLATED) as target:
        for original, payload in members:
            item = zipfile.ZipInfo(original.filename, date_time=(1980, 1, 1, 0, 0, 0))
            item.compress_type = zipfile.ZIP_DEFLATED
            item.external_attr = original.external_attr or (0o600 << 16)
            item.create_system = 0
            target.writestr(item, payload)
    os.replace(temp, path)


def write_pptx(route, output):
    output = Path(output)
    output.parent.mkdir(parents=True, exist_ok=True)
    presentation = build_presentation(route)
    presentation.save(output)
    normalize_zip(output)


def main():
    if len(sys.argv) != 3:
        print("Usage: python render_pptx.py <tech-route.json> <output.pptx>")
        return 2
    write_pptx(load_route(sys.argv[1]), sys.argv[2])
    print(f"Wrote {sys.argv[2]}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
